"""Pruebas de la agenda (versión 2) con pytest."""

from openpyxl import Workbook

from src.contactos import (
    normalizar_nombre,
    clave_nombre,
    nombre_es_valido,
    telefono_es_valido,
    nombre_ya_existe,
    crear_agenda,
    anadir_contacto,
    buscar_por_nombre,
    buscar_por_telefono,
    eliminar_contacto,
    ordenar_contactos,
    mostrar_contactos,
    guardar_contactos,
    cargar_contactos,
)


def test_normalizar_nombre_usa_title_y_quita_espacios():
    assert normalizar_nombre("  ana perez ") == "Ana Perez"


def test_nombre_no_vale_vacio_espacios_o_none():
    assert nombre_es_valido("") is False
    assert nombre_es_valido("   ") is False
    assert nombre_es_valido(None) is False
    assert nombre_es_valido("Ana Perez") is True


def test_telefono_valido():
    assert telefono_es_valido("612345678") is True


def test_telefono_invalido():
    assert telefono_es_valido("12345678") is False  # 8 dígitos
    assert telefono_es_valido("6123456789") is False  # 10 dígitos
    assert telefono_es_valido("61a345678") is False  # letras
    assert telefono_es_valido("012345678") is False  # empieza por 0
    assert telefono_es_valido("612 45678") is False  # espacios
    assert telefono_es_valido("612-45678") is False  # guiones
    assert telefono_es_valido(None) is False


def test_anadir_contacto_ok_title():
    agenda = crear_agenda()
    assert anadir_contacto(agenda, "ana perez", "612345678") is True
    assert len(agenda) == 1
    assert agenda[0] == {"nombre": "Ana Perez", "telefono": "612345678"}


def test_anadir_no_permite_nombre_duplicado_sin_mayusculas():
    agenda = crear_agenda()
    anadir_contacto(agenda, "Ana Perez", "612345678")
    assert anadir_contacto(agenda, "ANA PEREZ", "623456789") is False
    assert len(agenda) == 1


def test_anadir_si_permite_telefono_duplicado():
    agenda = crear_agenda()
    anadir_contacto(agenda, "Ana Perez", "612345678")
    assert anadir_contacto(agenda, "Luis", "612345678") is True
    assert len(agenda) == 2


def test_anadir_rechaza_datos_invalidos():
    agenda = crear_agenda()
    assert anadir_contacto(agenda, "   ", "612345678") is False
    assert anadir_contacto(agenda, "Ana", "123") is False
    assert len(agenda) == 0


def test_nombre_ya_existe():
    agenda = crear_agenda()
    anadir_contacto(agenda, "Ana Perez", "612345678")
    assert nombre_ya_existe(agenda, "ana perez") is True
    assert nombre_ya_existe(agenda, "Luis") is False


def test_clave_nombre_ignora_mayusculas_y_espacios():
    assert clave_nombre("  ANA  PEREZ ") == "ana perez"


def test_buscar_por_nombre_exacta_sin_mayusculas():
    agenda = crear_agenda()
    anadir_contacto(agenda, "Ana Perez", "612345678")
    assert len(buscar_por_nombre(agenda, "ANA PEREZ")) == 1
    assert buscar_por_nombre(agenda, "ana") == []
    assert buscar_por_nombre(agenda, "Luis") == []


def test_buscar_por_telefono_exacta():
    agenda = crear_agenda()
    anadir_contacto(agenda, "Ana Perez", "612345678")
    anadir_contacto(agenda, "Luis", "612345678")
    assert len(buscar_por_telefono(agenda, "612345678")) == 2
    assert buscar_por_telefono(agenda, "699999999") == []


def test_eliminar_sin_confirmacion():
    agenda = crear_agenda()
    anadir_contacto(agenda, "Ana Perez", "612345678")
    assert eliminar_contacto(agenda, "ana perez") is True
    assert len(agenda) == 0
    assert eliminar_contacto(agenda, "Ana Perez") is False


def test_mostrar_contactos():
    agenda = crear_agenda()
    assert mostrar_contactos(agenda) == "Agenda vacia."
    anadir_contacto(agenda, "Ana Perez", "612345678")
    assert mostrar_contactos(agenda) == "Ana Perez;612345678"


def test_ordenar_az_y_za():
    agenda = crear_agenda()
    anadir_contacto(agenda, "Luis", "623456789")
    anadir_contacto(agenda, "ana perez", "612345678")
    anadir_contacto(agenda, "Bea", "634567890")
    ordenados = ordenar_contactos(agenda, "az")
    assert [c["nombre"] for c in ordenados] == ["Ana Perez", "Bea", "Luis"]
    inversos = ordenar_contactos(agenda, "za")
    assert [c["nombre"] for c in inversos] == ["Luis", "Bea", "Ana Perez"]


def test_mostrar_az_y_za():
    agenda = crear_agenda()
    anadir_contacto(agenda, "Luis", "623456789")
    anadir_contacto(agenda, "ana perez", "612345678")
    assert mostrar_contactos(agenda, "az") == "Ana Perez;612345678\nLuis;623456789"
    assert mostrar_contactos(agenda, "za") == "Luis;623456789\nAna Perez;612345678"


def test_guardar_y_cargar_xlsx(tmp_path):
    ruta = str(tmp_path / "contactos.xlsx")
    agenda = crear_agenda()
    anadir_contacto(agenda, "Ana Perez", "612345678")
    guardar_contactos(agenda, ruta)
    cargada = cargar_contactos(ruta)
    assert cargada == agenda


def test_cargar_sin_fichero_crea_con_cabecera(tmp_path):
    from openpyxl import load_workbook
    ruta = str(tmp_path / "nuevo.xlsx")
    agenda = cargar_contactos(ruta)
    assert agenda == []
    libro = load_workbook(ruta)
    hoja = libro.active
    assert hoja["A1"].value == "Nombre"
    assert hoja["B1"].value == "Número"


def test_cargar_ignora_filas_rotas_y_duplicados(tmp_path):
    ruta = str(tmp_path / "contactos.xlsx")
    libro = Workbook()
    hoja = libro.active
    hoja["A1"] = "Nombre"
    hoja["B1"] = "Número"
    hoja.append(["Ana Perez", "612345678"])
    hoja.append(["Solo nombre", None])
    hoja.append(["ANA PEREZ", "623456789"])
    hoja.append(["Luis", "maltelefono"])
    hoja.append([None, None])
    libro.save(ruta)
    agenda = cargar_contactos(ruta)
    assert agenda == [{"nombre": "Ana Perez", "telefono": "612345678"}]
