"""Logica de la agenda telefonica (version 2).

Guarda los contactos como una lista de diccionarios:
[{"nombre": "Ana Perez", "telefono": "612345678"}]

No usa clases, solo funciones sencillas.
"""

from openpyxl import Workbook, load_workbook

RUTA_DATOS = "data/contactos.xlsx"
CABECERA_NOMBRE = "Nombre"
CABECERA_NUMERO = "Número"


def normalizar_nombre(nombre):
    """Quita espacios de fuera y usa title(). Ej: ana perez -> Ana Perez."""
    if nombre is None:
        return ""
    return " ".join(nombre.strip().split()).title()


def clave_nombre(nombre):
    """Clave para comparar sin distinguir mayusculas."""
    if nombre is None:
        return ""
    return " ".join(str(nombre).strip().split()).lower()


def nombre_es_valido(nombre):
    """Dice si el nombre vale: no vacio ni solo espacios."""
    if nombre is None:
        return False
    if nombre.strip() == "":
        return False
    return True


def telefono_es_valido(telefono):
    """Dice si el telefono vale: 9 digitos, sin empezar por 0."""
    if telefono is None:
        return False
    texto = str(telefono).strip()
    if len(texto) != 9:
        return False
    if not texto.isdigit():
        return False
    if texto.startswith("0"):
        return False
    return True


def nombre_ya_existe(agenda, nombre):
    """Dice si ese nombre ya esta en la agenda (sin mayusculas)."""
    buscado = clave_nombre(nombre)
    for contacto in agenda:
        if clave_nombre(contacto["nombre"]) == buscado:
            return True
    return False


def crear_agenda():
    """Devuelve una agenda vacia."""
    return []


def anadir_contacto(agenda, nombre, telefono):
    """Anade un contacto si los datos valen y el nombre no existe.

    Devuelve True si lo anade, False si no.
    Los telefonos duplicados si se permiten, los nombres no.
    """
    if not nombre_es_valido(nombre):
        return False
    if not telefono_es_valido(telefono):
        return False
    nombre_final = normalizar_nombre(nombre)
    telefono_final = str(telefono).strip()
    if nombre_ya_existe(agenda, nombre_final):
        return False
    agenda.append({"nombre": nombre_final, "telefono": telefono_final})
    return True


def buscar_por_nombre(agenda, nombre):
    """Busca por nombre exacto (da igual mayusculas). Devuelve lista."""
    buscado = clave_nombre(nombre)
    encontrados = []
    for contacto in agenda:
        if clave_nombre(contacto["nombre"]) == buscado:
            encontrados.append(contacto)
    return encontrados


def buscar_por_telefono(agenda, telefono):
    """Busca por telefono exacto. Devuelve lista."""
    if telefono is None:
        return []
    buscado = str(telefono).strip()
    encontrados = []
    for contacto in agenda:
        if contacto["telefono"] == buscado:
            encontrados.append(contacto)
    return encontrados


def eliminar_contacto(agenda, nombre):
    """Borra el contacto con ese nombre. Sin pedir confirmacion.

    Devuelve True si lo borro, False si no existia.
    """
    buscado = clave_nombre(nombre)
    for i in range(len(agenda)):
        if clave_nombre(agenda[i]["nombre"]) == buscado:
            del agenda[i]
            return True
    return False


def ordenar_contactos(agenda, sentido="az"):
    """Devuelve copia ordenada por nombre (sin mayusculas). az o za."""
    copia = list(agenda)
    copia.sort(key=lambda c: clave_nombre(c["nombre"]))
    if sentido == "za":
        copia.reverse()
    return copia


def mostrar_contactos(agenda, orden="llegada"):
    """Devuelve el texto con todos los contactos para imprimir."""
    if len(agenda) == 0:
        return "Agenda vacia."
    if orden == "az":
        agenda = ordenar_contactos(agenda, "az")
    elif orden == "za":
        agenda = ordenar_contactos(agenda, "za")
    lineas = []
    for contacto in agenda:
        lineas.append(contacto["nombre"] + ";" + contacto["telefono"])
    texto = ""
    for i in range(len(lineas)):
        texto = texto + lineas[i]
        if i < len(lineas) - 1:
            texto = texto + chr(10)
    return texto


def guardar_contactos(agenda, ruta=RUTA_DATOS):
    """Guarda toda la agenda en el xlsx con cabecera Nombre | Numero."""
    libro = Workbook()
    hoja = libro.active
    hoja["A1"] = CABECERA_NOMBRE
    hoja["B1"] = CABECERA_NUMERO
    fila = 2
    for contacto in agenda:
        hoja.cell(row=fila, column=1, value=contacto["nombre"])
        hoja.cell(row=fila, column=2, value=contacto["telefono"])
        fila += 1
    libro.save(ruta)


def cargar_contactos(ruta=RUTA_DATOS):
    """Lee el xlsx y devuelve la agenda.

    Si no existe, crea uno con cabecera y devuelve lista vacia.
    Las filas vacias o rotas se ignoran. Con nombres duplicados,
    se queda con el primero.
    """
    agenda = []
    try:
        libro = load_workbook(ruta)
    except FileNotFoundError:
        guardar_contactos(agenda, ruta)
        return agenda
    hoja = libro.active
    primera = True
    for fila in hoja.iter_rows(values_only=True):
        if primera:
            primera = False
            continue
        if fila is None:
            continue
        nombre = fila[0] if len(fila) > 0 else None
        telefono = fila[1] if len(fila) > 1 else None
        if nombre is None and telefono is None:
            continue
        anadir_contacto(agenda, nombre, telefono)
    return agenda
